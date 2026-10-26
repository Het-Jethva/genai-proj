"""Build the formatted CCA report from the verified project state."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "Northbridge_Guide_Report.docx"
EVAL_PATH = ROOT / "output" / "evaluation_results.json"
CORPUS_PATH = ROOT / "data" / "campus_records.jsonl"

NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x4B, 0x55, 0x65)
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"


def load_project_stats(report):
    records = [
        json.loads(line)
        for line in CORPUS_PATH.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    return {
        "record_count": len(records),
        "category_count": len({record["category"] for record in records}),
        "evaluation_count": report["total_cases"],
        "supported_count": report["supported_cases"],
        "refusal_count": report["refusal_cases"],
    }


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table, color="D6DCE5", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = GRAY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(header, after=2, line=1.0)
    run = header.add_run("NORTHBRIDGE GUIDE  /  CCA PROJECT")
    set_run_font(run, size=8.5, color=GRAY, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, before=2, after=0, line=1.0)
    run = footer.add_run("ECSDI24402  /  Page ")
    set_run_font(run, size=8.5, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_title_block(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=4, line=1.0)
    run = p.add_run("CCA PROJECT REPORT")
    set_run_font(run, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=4, line=1.0)
    run = p.add_run("Northbridge Guide")
    set_run_font(run, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=16, line=1.1)
    run = p.add_run("An evidence-first campus information assistant using retrieval-augmented generation")
    set_run_font(run, size=13.5, color=GRAY)

    meta = [
        ("Course", "Gen AI and applications [ECSDI24402]"),
        ("CCA choice", "Option 1, Project"),
        ("Team", "[Student 1] and [Student 2]"),
        ("Faculty", "[Faculty Name]"),
        ("Date", "31 August 2026"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    set_table_geometry(table, [1800, 7560])
    set_table_borders(table, color="D6DCE5", size="4")
    for row, (label, value) in zip(table.rows, meta):
        set_cell_shading(row.cells[0], PALE_BLUE)
        p1 = row.cells[0].paragraphs[0]
        set_paragraph_spacing(p1, after=0, line=1.0)
        r1 = p1.add_run(label)
        set_run_font(r1, size=9.5, color=NAVY, bold=True)
        p2 = row.cells[1].paragraphs[0]
        set_paragraph_spacing(p2, after=0, line=1.0)
        r2 = p2.add_run(value)
        set_run_font(r2, size=9.5, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_para(doc, text, *, bold_prefix=None, color=GRAY, after=6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.10)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, color=NAVY, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=5, line=1.10)
        r = p.add_run(item)
        set_run_font(r, color=GRAY)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=5, line=1.10)
        r = p.add_run(item)
        set_run_font(r, color=GRAY)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="CBD5E1", size="4")
    set_cell_shading(table.cell(0, 0), LIGHT_GRAY)
    p = table.cell(0, 0).paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.0)
    for index, line in enumerate(lines.splitlines()):
        if index:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=9, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="B7C7E1", size="6")
    set_cell_shading(table.cell(0, 0), PALE_BLUE)
    p = table.cell(0, 0).paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.10)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, color=NAVY, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_metrics_table(doc, report):
    add_heading(doc, "Evaluation results", 2)
    add_para(doc, "The offline suite runs without the hosted model so the baseline remains repeatable.")
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2500, 1300, 5560])
    set_table_borders(table)
    headers = ["Metric", "Result", "Meaning"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(value)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    rows = [
        ("Retrieval recall@5", f"{report['retrieval_recall_at_k']:.0%}", "Every supported case names at least one expected record in the top five."),
        ("Refusal accuracy", f"{report['refusal_accuracy']:.0%}", "Supported and out-of-scope cases received the expected decision."),
        ("Citation validity", f"{report['citation_validity']:.0%}", "Every supported fallback answer cites a record present in its evidence list."),
        ("Answer keyword coverage", f"{report['answer_keyword_coverage']:.1%}", "The fallback answer contains most keywords in each curated reference answer."),
    ]
    for metric, result, meaning in rows:
        cells = table.add_row().cells
        values = [metric, result, meaning]
        for index, (cell, value) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.05)
            r = p.add_run(value)
            set_run_font(r, size=9.3, color=NAVY if index == 1 else GRAY, bold=index == 1)


def add_rubric_table(doc):
    add_heading(doc, "Rubric mapping", 2)
    add_para(doc, "The implementation and documentation map directly to the 15-mark CCA criteria.")
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2500, 3900, 2960])
    set_table_borders(table)
    headers = ["Criterion", "Project evidence", "Presentation move"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(value)
        set_run_font(r, size=9.3, color=NAVY, bold=True)
    rows = [
        ("Understanding / 3", "RAG pipeline, evidence-first prompt, retrieval scores, refusal policy, dataset card.", "Explain the path from JSONL record to cited answer."),
        ("Quality / 3", "Schema validation, hybrid index, provider adapter, fallback, citation checking, tests.", "Run validation, unit tests, and the Evaluation tab."),
        ("Creativity / 3", "Fictional Northbridge corpus, evidence trail, refusal behavior, Corpus view.", "Show a supported question followed by the cafeteria-menu refusal."),
        ("Application / 4", "Nine campus areas covering academic, financial, library, career, club, facilities, scholarship, and wellbeing needs.", "Explain how a student gets a next step from one interface."),
        ("Documentation / 2", "Report, dataset card, methodology, demo script, rubric map, README, and PPTX.", "Keep the synthetic-data limitation visible."),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row_values)):
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.05)
            r = p.add_run(value)
            set_run_font(r, size=8.8, color=NAVY if index == 0 else GRAY, bold=index == 0)


def build():
    report = json.loads(EVAL_PATH.read_text(encoding="utf-8"))
    stats = load_project_stats(report)
    doc = Document()
    style_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "Project boundary",
        "Northbridge Institute of Technology, its policies, dates, amounts, and contact number are fictional. This report describes a classroom prototype, not an official university information system.",
    )

    add_heading(doc, "Abstract", 1)
    add_para(doc, "Northbridge Guide is a Streamlit prototype that answers campus-helpdesk questions from a self-prepared fictional dataset. The project demonstrates a complete retrieval-augmented generation pipeline. It validates JSONL records, retrieves relevant evidence, sends only that evidence to an OpenRouter-compatible model when configured, checks citations, and refuses questions that the corpus cannot support. A deterministic fallback keeps the demo and evaluation usable without a hosted key.")
    add_para(doc, f"The corpus contains {stats['record_count']} records across {stats['category_count']} campus areas. The offline evaluation contains {stats['evaluation_count']} cases, including {stats['refusal_count']} out-of-scope questions. The current baseline achieved {report['retrieval_recall_at_k']:.0%} retrieval recall@5, {report['refusal_accuracy']:.0%} refusal accuracy, {report['citation_validity']:.0%} citation validity, and {report['answer_keyword_coverage']:.1%} answer keyword coverage. These values describe this fictional corpus and should not be read as claims about a real university service.")

    add_heading(doc, "1. Problem statement", 1)
    add_para(doc, "Students often need small answers with exact details: a fee deadline, a borrowing limit, a form, a room-booking rule, or the next office to contact. A generic chatbot can answer quickly but may invent a date or policy when it lacks a source. This project tests a narrower design. The assistant should answer only when the prepared corpus provides enough evidence, show the records it used, and say when it cannot answer.")

    add_heading(doc, "2. Objectives", 1)
    add_numbered(doc, [
        "Prepare an original, inspectable dataset for a fictional campus.",
        "Build a working RAG application with a browser interface.",
        "Ground generated responses in retrieved records and display citations.",
        "Measure retrieval and refusal behavior with a repeatable test set.",
        "Document limitations and safe deployment practices.",
    ])

    add_heading(doc, "3. Dataset preparation", 1)
    add_para(doc, f"The team prepared {stats['record_count']} fictional records in data/campus_records.jsonl. The corpus covers academics, examinations, fees, library, scholarships, clubs and societies, placements and careers, facilities and IT, and safety and wellbeing. Each record has a stable ID, category, title, content, fictional source label, synthetic effective date, retrieval tags, confidence marker, and source_type value.")
    add_para(doc, "The dataset uses short, coherent records instead of long unstructured files. This makes each retrieved unit easy to inspect and cite. Records include both direct answers and boundaries. For example, the attendance record says that each course outline sets its own participation rule rather than inventing one institute-wide percentage. The emergency record includes a warning that its phone number is fictional.")
    add_para(doc, "The full data card is in docs/dataset_card.md. It states the intended classroom use, out-of-scope decisions, maintenance steps, and limitations.")

    add_heading(doc, "4. System design", 1)
    add_code_block(doc, "User question\n    |\n    v\nValidation and length limit\n    |\n    v\nHybrid retrieval: word score + character score + concept score\n    |\n    v\nAnswerability check and top-k evidence\n    |                         \\\n    |                          \\\n    v                           v\nOpenRouter/free generation      Deterministic grounded fallback\n    |                           /\n    v                          /\nCitation validation and refusal guard\n    |\n    v\nAnswer, evidence trail, score, provider status")

    add_heading(doc, "Retrieval", 2)
    add_para(doc, "The default index uses weighted word similarity, character n-gram similarity, and a small concept-expansion table. Word similarity handles direct terms. Character n-grams help with spelling variation and exact fragments. Concept groups connect phrases such as exam and examination, or fee and installment. The optional sentence-transformers path adds dense cosine similarity when explicitly enabled.")

    add_heading(doc, "Generation", 2)
    add_para(doc, "The hosted path sends a system instruction and an evidence block to the /api/v1/chat/completions endpoint. The configured model defaults to openrouter/free. The prompt asks the model to answer only from the evidence, avoid invented facts, cite labels such as [R1], and keep the response short. The provider adapter reads the API key from environment configuration and never writes it to the repository.")
    add_para(doc, "When no key is present or the request fails, the service returns the strongest retrieved record as a deterministic answer with a valid citation. The UI labels this provider path so the evaluator knows whether the hosted model ran.")

    add_heading(doc, "Guardrails", 2)
    add_bullets(doc, [
        "The app rejects blank questions and questions longer than 800 characters.",
        "The answer service checks retrieval score, direct query coverage, concept-group coverage, and unsupported-term ratio.",
        "Low-evidence questions receive a refusal and do not cite unrelated records.",
        "Model citations are limited to the retrieved evidence labels.",
        "The public demo remains open to visitors but limits each browser session to 30 requests.",
        "The app and report repeat the fictional-data disclaimer.",
    ])

    add_heading(doc, "5. User interface", 1)
    add_para(doc, "The Streamlit app has three views.")
    add_bullets(doc, [
        "Ask: question form, prepared prompts, answer, retrieval metadata, and evidence trail.",
        "Evaluation: repeatable offline metrics and per-case results.",
        "Corpus: filters and readable source records.",
    ])
    add_para(doc, "The interface uses a cool near-white surface, a single cobalt signal color, technical typography, and a graphite evidence surface. It keeps the main action left aligned and uses borders instead of decorative gradients or card stacks. The full CSS includes responsive rules for narrow screens and visible focus states.")

    add_heading(doc, "6. Evaluation", 1)
    add_para(doc, f"The test set contains {stats['evaluation_count']} cases. {stats['supported_count']} are expected to be answerable, and {stats['refusal_count']} are expected to be refused because the corpus has no current rector record, cafeteria menu, or placement-package outcome data.")
    add_metrics_table(doc, report)
    add_para(doc, "The command writes the machine-readable result to output/evaluation_results.json. The evaluation intentionally runs without the hosted model. Model answers can vary with provider availability, so the offline baseline provides a stable reference point.")

    add_heading(doc, "7. Responsible use and limitations", 1)
    add_para(doc, "This is a classroom prototype over fictional data. It is not an official university information system. The synthetic dates, charges, contact number, and rules must not guide real student decisions. A real version would require approved source documents, a content owner for each policy area, expiry monitoring, privacy review, access controls, multilingual testing, and human review for high-impact answers.")
    add_para(doc, "The hosted free-model route also depends on external availability and quota. Public deployment must keep the key in the host secrets manager. The request cap lowers risk but does not make a public key unlimited or abuse-proof.")

    add_heading(doc, "8. Reproduction steps", 1)
    add_code_block(doc, "uv sync --dev\nuv run python scripts\\validate_dataset.py\nuv run python -m unittest discover -s tests -v\nuv run python scripts\\evaluate.py\nuv run streamlit run app.py")
    add_para(doc, "For hosted generation, set OPENROUTER_API_KEY and keep OPENROUTER_MODEL=openrouter/free. Without a key, the app runs through the grounded fallback.")

    add_heading(doc, "9. Rubric mapping", 1)
    add_rubric_table(doc)

    add_heading(doc, "10. Conclusion", 1)
    add_para(doc, "Northbridge Guide shows the whole GenAI workflow in a small project: prepare data, retrieve evidence, generate a response, show citations, refuse unsupported questions, and measure the result. Its most important design choice is the refusal path. The prototype does not pretend that a language model can answer every campus question. It makes the boundary visible, testable, and easy to explain during the CCA presentation.")

    add_heading(doc, "References", 1)
    add_para(doc, "OpenRouter, Quickstart Guide, https://openrouter.ai/docs/quickstart", after=4)
    add_para(doc, "OpenRouter, Free Models Router, https://openrouter.ai/docs/cookbook/get-started/free-models-router-playground", after=4)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Northbridge Guide CCA Project Report"
    doc.core_properties.subject = "ECSDI24402 Gen AI and applications"
    doc.core_properties.author = "[Student 1] and [Student 2]"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
